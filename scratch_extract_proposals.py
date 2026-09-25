import os
import json
from pypdf import PdfReader
from pptx import Presentation

proposal_dir = 'Proposals'
results = {}

for fname in sorted(os.listdir(proposal_dir)):
    fpath = os.path.join(proposal_dir, fname)
    if fname.endswith('.pdf'):
        reader = PdfReader(fpath)
        pages_text = []
        for i, page in enumerate(reader.pages):
            t = page.extract_text() or ''
            pages_text.append({'page': i + 1, 'text': t})
        results[fname] = {'type': 'pdf', 'pages_count': len(reader.pages), 'pages': pages_text}
    elif fname.endswith('.pptx'):
        prs = Presentation(fpath)
        slides_text = []
        for i, slide in enumerate(prs.slides):
            slide_elements = []
            for shape in slide.shapes:
                if shape.has_text_frame:
                    txt = '\n'.join([p.text.strip() for p in shape.text_frame.paragraphs if p.text.strip()])
                    if txt:
                        slide_elements.append({'kind': 'text', 'content': txt})
                elif shape.has_table:
                    table_rows = []
                    for row in shape.table.rows:
                        row_vals = [cell.text.strip() for cell in row.cells]
                        table_rows.append(row_vals)
                    slide_elements.append({'kind': 'table', 'rows': table_rows})
            slides_text.append({'slide': i + 1, 'elements': slide_elements})
        results[fname] = {'type': 'pptx', 'slides_count': len(prs.slides), 'slides': slides_text}

with open('scratch_proposals_extracted.json', 'w', encoding='utf-8') as f:
    json.dump(results, f, indent=2, ensure_ascii=False)

print('Extracted successfully!')
